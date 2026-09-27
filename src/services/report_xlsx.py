from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config.settings import CLASSES

NIGHT, PLUM, LILAC, SILVER = "1D1526", "5C4B73", "D4C4E8", "EAE6F2"
BAND = "F5F1FA"
LINE = "DCD2EA"
GOOD, GOOD_BG = "1E6B3A", "E3F4E8"
BAD, BAD_BG = "9B1C1C", "FBE3E3"
WARN, WARN_BG = "8A5A00", "FFF3D6"

F_TITLE = Font(name="Segoe UI", size=18, bold=True, color=SILVER)
F_SUB = Font(name="Segoe UI", size=10, color=LILAC)
F_HEAD = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="Segoe UI", size=10, color=NIGHT)
F_BOLD = Font(name="Segoe UI", size=10, bold=True, color=NIGHT)
F_SECTION = Font(name="Segoe UI", size=12, bold=True, color=PLUM)
FILL_NIGHT = PatternFill("solid", fgColor=NIGHT)
FILL_HEAD = PatternFill("solid", fgColor=PLUM)
FILL_BAND = PatternFill("solid", fgColor=BAND)
THIN = Border(bottom=Side(style="thin", color=LINE))
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _banner(ws, title: str, subtitle: str, width: int):
    for r in (1, 2):
        for c in range(1, max(width, 6) + 1):
            ws.cell(r, c).fill = FILL_NIGHT
    ws.cell(1, 1, title).font = F_TITLE
    ws.cell(2, 1, subtitle).font = F_SUB
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18


def _table(ws, df: pd.DataFrame, top: int, left: int = 1, widths: dict | None = None):
    for j, col in enumerate(df.columns):
        c = ws.cell(top, left + j, str(col))
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, CENTER
    ws.row_dimensions[top].height = 32
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(top + i, left + j, v)
            c.font, c.border = F_BODY, THIN
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if i % 2 == 0:
                c.fill = FILL_BAND
    for j, col in enumerate(df.columns):
        w = (widths or {}).get(col)
        if w is None:
            longest = max([len(str(col))] + [len(str(x)) for x in df[col].head(200).tolist()])
            w = min(max(10, longest + 2), 48)
        ws.column_dimensions[get_column_letter(left + j)].width = w
    return top + len(df) + 2


def _mark(ws, df: pd.DataFrame, top: int, col: str, good=("Yes", "Correct"), bad=("No", "Incorrect")):
    if col not in df.columns:
        return
    j = list(df.columns).index(col) + 1
    for i, v in enumerate(df[col].tolist(), start=1):
        c = ws.cell(top + i, j)
        if v in good:
            c.font, c.fill = Font(name="Segoe UI", size=10, bold=True, color=GOOD), _fill(GOOD_BG)
        elif v in bad:
            c.font, c.fill = Font(name="Segoe UI", size=10, bold=True, color=BAD), _fill(BAD_BG)


def style_sheet(ws, title: str | None = None):
    for c in ws[1]:
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, CENTER
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.font, c.border = F_BODY, THIN
            if r % 2 == 1:
                c.fill = FILL_BAND
    for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, longest + 2), 48)
    ws.freeze_panes = "B2"
    ws.sheet_properties.tabColor = PLUM


def _pct(v):
    return None if v is None else round(float(v) * 100, 1)


def _confusion(df: pd.DataFrame, pred_col: str) -> pd.DataFrame:
    lab = df[df["Actual class"] != ""]
    m = pd.DataFrame(0, index=CLASSES, columns=CLASSES)
    for a, p in zip(lab["Actual class"], lab[pred_col]):
        if a in m.index and p in m.columns:
            m.loc[a, p] += 1
    m.index.name = "Actual \\ Predicted"
    return m.reset_index()


def comparison_workbook(df: pd.DataFrame, s: dict, meta: dict) -> io.BytesIO:
    wb = Workbook()
    stamp = datetime.now().strftime("%d %b %Y %H:%M")

    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = NIGHT
    _banner(ws, "SonicSentinel AI – Model Prediction & Confidence Comparison",
            f"SRS deliverable 6 · unseen TEST recordings · generated {stamp}", 8)
    r = 4
    ws.cell(r, 1, "Overall comparison summary").font = F_SECTION
    r += 1
    counts = s.get("per_class_counts", {})
    rows = [
        ("Recordings in the report", s.get("with_ground_truth"), ""),
        ("SRS size requirement (≥ 100 recordings, ≥ 10 per class)",
         "Met" if s.get("meets_srs_size") else "Not met",
         f"smallest class: {min(counts.values()) if counts else 0} recordings"),
        ("Python model", meta.get("python_version", ""), ""),
        ("GTM model", meta.get("gtm_version", ""), ""),
        ("Python accuracy (%)", _pct(s.get("python_accuracy")), ""),
        ("GTM accuracy (%)", _pct(s.get("gtm_accuracy")), ""),
        ("Final (combined) accuracy (%)", _pct(s.get("final_accuracy")), "decision after comparison, fusion and rules"),
        ("Model agreement (%)", _pct(s.get("agreement_rate")), "both models chose the same class"),
        ("Disagreements", s.get("disagreements"),
         f"Python right {s.get('disagreements_python_right', 0)} · GTM right {s.get('disagreements_gtm_right', 0)}"),
        ("Mean top-class confidence difference (%)", _pct(s.get("mean_confidence_difference")), "|Python − GTM|"),
        ("Sent to manual review (%)", _pct(s.get("manual_review_rate")), ""),
        ("Automatic alerts generated", s.get("alerts_generated"), ""),
    ]
    r = _table(ws, pd.DataFrame(rows, columns=["Metric", "Value", "Note"]), r,
               widths={"Metric": 52, "Value": 26, "Note": 48})
    ok_row = 5 + 2
    ws.cell(ok_row, 2).font = Font(name="Segoe UI", size=10, bold=True,
                                   color=GOOD if s.get("meets_srs_size") else BAD)
    ws.cell(r, 1, "How to read this report").font = F_SECTION
    notes = [
        "Comparison – one row per recording with every column required by the SRS (both models' confidence for every class).",
        "Per class – recordings, Python / GTM / final accuracy and agreement for each class.",
        "Confusion matrices – rows are the actual class, columns the predicted class (diagonal = correct).",
        "Disagreements – every recording where the two models chose different classes, with an explanation.",
        "Final decision: if both models agree their scores are combined as independent evidence; if they disagree "
        f"the Python model has weight {meta.get('python_weight', '')}. Critical classes still need confirmation by the alert rules.",
    ]
    for n in notes:
        r += 1
        c = ws.cell(r, 1, "• " + n)
        c.font, c.alignment = F_BODY, WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 30

    wc = wb.create_sheet("Comparison")
    wc.sheet_properties.tabColor = PLUM
    _banner(wc, "Per-recording comparison", "Actual class vs both models, final decision and result", len(df.columns))
    top = 4
    _table(wc, df.drop(columns=["Correct"], errors="ignore"), top,
           widths={"Explanation of disagreement": 70, "Filename": 30, "Audio ID": 24})
    cols = list(df.drop(columns=["Correct"], errors="ignore").columns)
    shown = df.drop(columns=["Correct"], errors="ignore")
    for col in ("Result", "Python correct", "GTM correct", "Class match"):
        _mark(wc, shown, top, col)
    for i, rec in enumerate(shown.to_dict("records"), start=1):
        for prefix, key in (("Python conf: ", "Python predicted class"), ("GTM conf: ", "GTM predicted class")):
            col = prefix + str(rec.get(key))
            if col in cols:
                c = wc.cell(top + i, cols.index(col) + 1)
                c.font, c.fill = F_BOLD, _fill(LILAC)
    for name in cols:
        if name.startswith(("Python conf:", "GTM conf:")) or "margin" in name or "difference" in name or name == "Final confidence":
            j = cols.index(name) + 1
            wc.column_dimensions[get_column_letter(j)].width = 13
            for i in range(1, len(shown) + 1):
                wc.cell(top + i, j).number_format = "0.00%"
    wc.freeze_panes = wc.cell(top + 1, 3)
    if cols:
        wc.auto_filter.ref = f"A{top}:{get_column_letter(len(cols))}{top + len(shown)}"
    else:
        wc.cell(top, 1, "No evaluated recordings yet – run Admin → Model comparison → Run test-set evaluation.").font = F_BODY

    wp = wb.create_sheet("Per class")
    wp.sheet_properties.tabColor = PLUM
    _banner(wp, "Per-class results", "Accuracy of each model and of the final decision for every class", 6)
    pc = s.get("per_class", {})
    pdf = pd.DataFrame([{"Class": c, "Recordings": v["recordings"],
                         "Python accuracy (%)": _pct(v["python_accuracy"]), "GTM accuracy (%)": _pct(v["gtm_accuracy"]),
                         "Final accuracy (%)": _pct(v["final_accuracy"]), "Agreement (%)": _pct(v["agreement"])}
                        for c, v in pc.items()])
    if not pdf.empty:
        _table(wp, pdf, 4, widths={"Class": 26})

    wm = wb.create_sheet("Confusion matrices")
    wm.sheet_properties.tabColor = PLUM
    _banner(wm, "Confusion matrices", "Rows = actual class · columns = predicted class · diagonal = correct", len(CLASSES) + 1)
    r = 4
    for title, col in (("Python model", "Python predicted class"), ("GTM model", "GTM predicted class"),
                       ("Final decision", "Final decision")):
        wm.cell(r, 1, title).font = F_SECTION
        if df.empty:
            wm.cell(r + 1, 1, "No data yet.").font = F_BODY
            r += 3
            continue
        m = _confusion(df, col)
        start = r + 1
        r = _table(wm, m, start, widths={"Actual \\ Predicted": 26, **{c: 11 for c in CLASSES}})
        mx = max(1, int(m[CLASSES].to_numpy().max()))
        for i in range(len(CLASSES)):
            for j in range(len(CLASSES)):
                v = int(m.iloc[i][CLASSES[j]])
                cell = wm.cell(start + 1 + i, 2 + j)
                cell.alignment = CENTER
                if v:
                    t = v / mx
                    rgb = [int(a + (b - a) * t) for a, b in ((0xF1, 0x5C), (0xEA, 0x4B), (0xF8, 0x73))]
                    cell.fill = _fill("".join(f"{x:02X}" for x in rgb))
                    cell.font = Font(name="Segoe UI", size=10, bold=i == j, color="FFFFFF" if t > 0.5 else NIGHT)
                if i == j:
                    cell.border = Border(left=Side(style="medium", color=PLUM), right=Side(style="medium", color=PLUM),
                                         top=Side(style="medium", color=PLUM), bottom=Side(style="medium", color=PLUM))

    wd = wb.create_sheet("Disagreements")
    wd.sheet_properties.tabColor = PLUM
    _banner(wd, "Major disagreements", "Different classes, or the same class with a confidence gap ≥ 0.40", 8)
    dd = df[df["Explanation of disagreement"] != ""][
        ["Audio ID", "Actual class", "Python predicted class", "GTM predicted class", "Top-class confidence difference",
         "Final decision", "Result", "Explanation of disagreement"]] if not df.empty else pd.DataFrame()
    if not dd.empty:
        _table(wd, dd, 4, widths={"Explanation of disagreement": 90, "Audio ID": 24})
        _mark(wd, dd, 4, "Result")
    else:
        wd.cell(4, 1, "No major disagreements.").font = F_BODY

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
